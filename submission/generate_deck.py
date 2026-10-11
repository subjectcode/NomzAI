import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color tokens
    C_BG_DARK = RGBColor(15, 23, 42)        # #0F172A deep slate
    C_BG_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC
    C_CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
    C_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0
    C_TEAL = RGBColor(13, 148, 136)         # #0D9488 Nomz Brand
    C_TEAL_DARK = RGBColor(15, 118, 110)    # #0F766E
    C_EMERALD = RGBColor(16, 185, 129)      # #10B981
    C_AMBER = RGBColor(245, 158, 11)        # #F59E0B
    C_TEXT_DARK = RGBColor(15, 23, 42)      # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    C_WHITE = RGBColor(255, 255, 255)

    base_dir = r"d:\Projects\Bob\Nomz"
    ss_dir = os.path.join(base_dir, "submission", "screenshots")

    def add_header(slide, title, category="NOMZ AI FOOD RESCUE ASSISTANT", is_dark=False):
        # Category pill/label
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_EMERALD if is_dark else C_TEAL

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = C_WHITE if is_dark else C_TEXT_DARK

    def set_slide_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def create_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, C_BG_DARK)

    # Accent decorative block
    accent = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(4.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = C_EMERALD
    accent.line.fill.background()

    # Title box
    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(11), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "HACKATHON SUBMISSION — HEALTHCARE & WELLBEING"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = C_EMERALD
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Nomz AI Food Rescue Assistant"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "Solusi Cerdas Penyelamat Bahan Makanan Dapur:\nPindai Bahan Sisa, Cegah Food Waste, dan Temukan Rekomendasi Masakan Sehat"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(36)

    p3 = tf.add_paragraph()
    p3.text = "Pengembang: Tim Nomz AI  |  Tech Stack: React Native, FastAPI, Google Gemini, Langflow V2, IBM Bob/MCP\nStatus: Produksi Siap Pakai (Universal Android APK v1.0.0 & Railway Backend Live)"
    p3.font.size = Pt(12)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: PROBLEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, C_BG_LIGHT)
    add_header(s2, "Tantangan Nyata: Food Waste Rumah Tangga & Dilema Dapur Anak Kos")

    # 3 Problem Cards
    cards_data = [
        ("Dilema 'Masak Apa Hari Ini?'", "Banyak orang, khususnya anak kos dan keluarga muda, bingung mengolah bahan mentah atau sisa makanan yang tersisa sedikit di kulkas sehingga akhirnya dibiarkan membusuk.", C_TEAL),
        ("Bahan Makanan Cepat Rusak", "Bahan segar seperti sayuran, cabai, tahu, dan daging mudah busuk dalam 2-4 hari jika tidak segera diolah. Pembelian berulang menambah beban finansial bulanan.", C_AMBER),
        ("Kesenjangan Resep Konvensional", "Aplikasi resep biasa menuntut bahan lengkap dari awal. Saat pengguna kekurangan 1-2 bahan, mereka batal memasak dan memilih beli makanan luar yang lebih boros.", RGBColor(239, 68, 68))
    ]

    for idx, (title, desc, accent_col) in enumerate(cards_data):
        cx = 0.8 + idx * 3.95
        card = create_card(s2, cx, 1.8, 3.75, 4.8)
        
        # Color bar top
        bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx + 0.3), Inches(2.1), Inches(0.8), Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent_col
        bar.line.fill.background()

        tb = s2.shapes.add_textbox(Inches(cx + 0.3), Inches(2.4), Inches(3.15), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(14)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = C_TEXT_MUTED
        p_desc.line_spacing = 1.3

    # ==========================================
    # SLIDE 3: SOLUTION
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, C_BG_LIGHT)
    add_header(s3, "Solusi Nomz: Dari Bahan Sisa Menjadi Hidangan Lezat & Bergizi")

    sol_cards = [
        ("1. Multimodal AI Scan", "Cukup foto isi kulkas atau meja dapur. AI mendeteksi bahan mentah (sayur, protein, bumbu) secara otomatis dengan confidence score visual."),
        ("2. Validasi Transparan", "Prinsip Human-in-the-Loop: Pengguna bebas menyesuaikan jumlah takaran, menghapus salah deteksi, atau menambah bumbu yang ada di dapur."),
        ("3. Rekomendasi Langflow V2", "Mesin orkestrasi mencocokkan resep masakan Nusantara berdasarkan stok yang ada, meminimalkan bahan tambahan yang harus dibeli."),
        ("4. Panduan Step-by-Step", "Instruksi memasak ringkas, estimasi biaya, takaran bumbu, dan kurasi video YouTube agar pengguna sukses mengeksekusi resep.")
    ]

    for idx, (title, desc) in enumerate(sol_cards):
        cx = 0.8 + idx * 2.95
        create_card(s3, cx, 1.8, 2.8, 4.8)
        
        badge = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.3), Inches(2.1), Inches(0.55), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_TEAL
        badge.line.fill.background()
        
        tb_num = s3.shapes.add_textbox(Inches(cx + 0.3), Inches(2.1), Inches(0.55), Inches(0.55))
        tf_num = tb_num.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.CENTER
        p_num.text = str(idx + 1)
        p_num.font.bold = True
        p_num.font.size = Pt(16)
        p_num.font.color.rgb = C_WHITE

        tb = s3.shapes.add_textbox(Inches(cx + 0.3), Inches(2.8), Inches(2.2), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.line_spacing = 1.3

    # ==========================================
    # SLIDE 4: PRODUCT DEMO (SCREENSHOTS)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, C_BG_LIGHT)
    add_header(s4, "Alur Produk: 5 Tahap Pengalaman Pengguna di Aplikasi Android")

    screens = [
        ("01-beranda.png", "1. Beranda", "Aksi cepat & metrik"),
        ("02-deteksi-bahan.png", "2. Deteksi AI", "Foto & kurasi bahan"),
        ("03-rekomendasi-resep.png", "3. Rekomendasi", "Pantry match scoring"),
        ("04-panduan-memasak.png", "4. Detail Resep", "Panduan 9 langkah"),
        ("05-profil-dampak.png", "5. Profil Dampak", "Penyelamatan pangan")
    ]

    for idx, (fname, title, subtitle) in enumerate(screens):
        fpath = os.path.join(ss_dir, fname)
        cx = 0.8 + idx * 2.4
        
        # Insert image
        if os.path.exists(fpath):
            s4.shapes.add_picture(fpath, Inches(cx), Inches(1.8), Inches(2.2), Inches(4.7))
        
        # Caption box
        tb = s4.shapes.add_textbox(Inches(cx), Inches(6.55), Inches(2.2), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK

        p_sub = tf.add_paragraph()
        p_sub.alignment = PP_ALIGN.CENTER
        p_sub.text = subtitle
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 5: AI WORKFLOW
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, C_BG_DARK)
    add_header(s5, "Arsitektur AI: Pipeline Multimodal & Orkestrasi Langflow V2", is_dark=True)

    wf_steps = [
        ("Tahap 1: Multimodal Vision", "Google Gemini 2.5/Flash", "Menganalisis foto bahan mentah. Mengembalikan JSON terstruktur nama bahan, estimasi kuantitas cerdas, dan confidence score (0.0 - 1.0)."),
        ("Tahap 2: Validasi Pengguna", "Human-in-the-Loop", "Pengguna mengonfirmasi foto dan hasil deteksi. Mengedit jumlah porsi via stepper, menambah atau menghapus bahan sebelum resep diminta."),
        ("Tahap 3: Mesin Langflow V2", "Multi-Agent Flow Orchestration", "Prompt chaining & retrieval berbasis kecocokan stok dapur (Pantry Match Score). Mengelompokkan bahan utama vs bumbu pelengkap umum."),
        ("Tahap 4: Kurasi Masakan", "FastAPI Response Engine", "Menghasilkan resep terstruktur: estimasi waktu, tingkat kesulitan, takaran bumbu, instruksi langkah demi langkah, dan video tutorial.")
    ]

    for idx, (title, tech, desc) in enumerate(wf_steps):
        top_pos = 1.8 + idx * 1.25
        card = create_card(s5, 0.8, top_pos, 11.73, 1.1, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
        
        # Number badge
        nb = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(top_pos + 0.25), Inches(0.6), Inches(0.6))
        nb.fill.solid()
        nb.fill.fore_color.rgb = C_EMERALD
        nb.line.fill.background()
        
        tb_n = s5.shapes.add_textbox(Inches(1.1), Inches(top_pos + 0.25), Inches(0.6), Inches(0.6))
        p_n = tb_n.text_frame.paragraphs[0]
        p_n.alignment = PP_ALIGN.CENTER
        p_n.text = str(idx + 1)
        p_n.font.bold = True
        p_n.font.size = Pt(16)
        p_n.font.color.rgb = C_WHITE

        # Text
        tb = s5.shapes.add_textbox(Inches(1.9), Inches(top_pos + 0.15), Inches(10.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{title}  —  "
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        
        run_tech = p.add_run()
        run_tech.text = f"[{tech}]"
        run_tech.font.size = Pt(12)
        run_tech.font.bold = True
        run_tech.font.color.rgb = C_EMERALD

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = RGBColor(203, 213, 225)

    # ==========================================
    # SLIDE 6: IBM BOB & MCP
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, C_BG_LIGHT)
    add_header(s6, "Peran IBM Bob & Model Context Protocol (MCP) dalam Siklus Rekayasa")

    # Left Card: Klarifikasi Integrasi
    create_card(s6, 0.8, 1.8, 5.65, 4.8)
    tb_left = s6.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.05), Inches(4.2))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "Peran Nyata: Developer & QA Orchestrator"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(12)

    bullets_l = [
        ("Bukan Chatbot Dalam Aplikasi", "IBM Bob tidak dipasang sebagai asisten chat UI bagi pengguna akhir, melainkan bertindak sebagai Agentic Development Environment."),
        ("MCP Tooling Integration", "Memanfaatkan server Model Context Protocol untuk membaca konteks basis kode, database schema, dan manajemen tugas secara otonom."),
        ("Automasi Regression Testing", "IBM Bob mengeksekusi skenario uji end-to-end, verifikasi schema JWT auth, integrasi Langflow V2, dan audit konsistensi APK build.")
    ]

    for title, desc in bullets_l:
        p_t = tf_l.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = C_TEAL_DARK
        
        r = p_t.add_run()
        r.text = desc
        r.font.bold = False
        r.font.color.rgb = C_TEXT_MUTED
        p_t.space_after = Pt(10)

    # Right Card: Nilai Tambah bagi Proyek
    create_card(s6, 6.85, 1.8, 5.65, 4.8)
    tb_right = s6.shapes.add_textbox(Inches(7.15), Inches(2.1), Inches(5.05), Inches(4.2))
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "Efisiensi Rekayasa & Keandalan Produksi"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(12)

    bullets_r = [
        ("Deteksi Dini Bug Arsitektur", "Mengidentifikasi kegagalan ABI x86 vs ARM pada build Android sebelumnya dan mengarahkan pembuatan Universal Release APK."),
        ("Validasi Kontrak API", "Memastikan keselarasan schema Pydantic antara backend FastAPI dengan DTO TypeScript di React Native secara deterministik."),
        ("Penyelarasan Prompt Langflow", "Menguji variasi prompt masakan Nusantara agar model AI tidak mengarang bahan non-eksisten atau langkah berbahaya.")
    ]

    for title, desc in bullets_r:
        p_t = tf_r.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = C_EMERALD
        
        r = p_t.add_run()
        r.text = desc
        r.font.bold = False
        r.font.color.rgb = C_TEXT_MUTED
        p_t.space_after = Pt(10)

    # ==========================================
    # SLIDE 7: TECH STACK & IMPLEMENTATION
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, C_BG_LIGHT)
    add_header(s7, "Teknologi & Implementasi: Stack Produksi Modern & Teruji")

    tech_boxes = [
        ("Mobile Frontend", "React Native 0.76 & Expo", "Antarmuka responsif 5-tab, Camera & PhotoPicker native, state management offline-ready, kompatibel Android 7.0+ (API 24-36)."),
        ("Backend Services", "FastAPI & Python 3.11", "Arsitektur RESTful asinkron, schema validation Pydantic V2, multipart image upload, JWT token & Google OAuth verification."),
        ("Database & Deploy", "PostgreSQL & Railway", "Penyimpanan relasional profil pengguna, histori resep tersimpan, dan metadata pantry. Deployed live di Railway Cloud."),
        ("Multimodal Vision", "Google Gemini AI", "Multimodal reasoning untuk identifikasi bahan mentah lokal (bawang merah, cabai rawit, tempe, daging) dengan JSON schema terstruktur."),
        ("Workflow Engine", "Langflow V2 Orchestration", "Mesin rekomendasi resep multi-agent dengan scoring kecocokan stok bahan (Pantry Match) dan fallback cerdas."),
        ("Media & Integrasi", "YouTube Data API v3", "Kurasi video tutorial memasak praktis yang relevan dengan resep yang disarankan langsung di dalam aplikasi.")
    ]

    for idx, (title, tech, desc) in enumerate(tech_boxes):
        col = idx % 3
        row = idx // 3
        cx = 0.8 + col * 3.95
        cy = 1.8 + row * 2.45
        
        create_card(s7, cx, cy, 3.75, 2.25)
        
        tb = s7.shapes.add_textbox(Inches(cx + 0.25), Inches(cy + 0.2), Inches(3.25), Inches(1.85))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        
        p_sub = tf.add_paragraph()
        p_sub.text = tech
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = C_TEAL
        p_sub.space_after = Pt(6)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = C_TEXT_MUTED
        p_desc.line_spacing = 1.25

    # ==========================================
    # SLIDE 8: TESTING & RESPONSIBLE AI
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8, C_BG_LIGHT)
    add_header(s8, "Keandalan Rekayasa & Penerapan Responsible AI")

    # Metrics Row
    m_data = [
        ("83 / 83", "Backend Test Passed (Pytest)", C_EMERALD),
        ("0 Error", "TypeScript Strict Compilation", C_TEAL),
        ("v1.0.0", "Release APK Universal 38.9 MB", C_AMBER),
        ("ARM + x86", "Kompatibilitas Multi-Arsitektur", C_TEXT_DARK)
    ]

    for idx, (val, lbl, col) in enumerate(m_data):
        cx = 0.8 + idx * 2.95
        card = create_card(s8, cx, 1.8, 2.8, 1.3)
        
        tb = s8.shapes.add_textbox(Inches(cx + 0.2), Inches(1.9), Inches(2.4), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = col
        
        p_l = tf.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(9.5)
        p_l.font.color.rgb = C_TEXT_MUTED

    # 3 Responsible AI Cards
    resp_cards = [
        ("Human-in-the-Loop Validation", "AI tidak pernah langsung mengeksekusi rekomendasi tanpa persetujuan pengguna. Pengguna memiliki kontrol penuh untuk memverifikasi bahan, mengubah kuantitas, atau menambah bahan manual."),
        ("Mitigasi Halusinasi & Kebersihan Data", "Prompt vision dilengkapi aturan ketat: hanya mendeteksi bahan yang kasat mata, melarang asumsi bumbu tak terlihat (seperti garam/minyak), dan menghasilkan confidence score transparan."),
        ("Batasan & Keselamatan Pangan", "Aplikasi tidak memberikan klaim medis atau diagnosis klinis. Pengguna diberi peringatan untuk memeriksa kesegaran bahan dan potensi alergi secara mandiri sebelum memasak.")
    ]

    for idx, (title, desc) in enumerate(resp_cards):
        cx = 0.8 + idx * 3.95
        create_card(s8, cx, 3.35, 3.75, 3.4)
        
        tb = s8.shapes.add_textbox(Inches(cx + 0.25), Inches(3.55), Inches(3.25), Inches(2.9))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(8)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = C_TEXT_MUTED
        p_desc.line_spacing = 1.3

    # ==========================================
    # SLIDE 9: IMPACT & SCALABILITY
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9, C_BG_LIGHT)
    add_header(s9, "Dampak Sosial, Kesejahteraan, & Peta Jalan Skalabilitas")

    # Left: Potensi Dampak
    create_card(s9, 0.8, 1.8, 5.65, 4.8)
    tb_imp = s9.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.05), Inches(4.2))
    tf_i = tb_imp.text_frame
    tf_i.word_wrap = True

    p = tf_i.paragraphs[0]
    p.text = "Potensi Dampak (Healthcare & Wellbeing)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(12)

    imp_items = [
        ("Pengurangan Food Waste Dapur", "Membantu menyelamatkan rata-rata 1-2 kg bahan makanan per rumah tangga setiap minggu yang sebelumnya berisiko terbuang."),
        ("Penghematan Finansial", "Menekan pengeluaran belanja makanan di luar dengan memaksimalkan konsumsi bahan yang sudah dibeli."),
        ("Peningkatan Kualitas Gizi", "Mendorong kebiasaan memasak makanan segar dan higienis sendiri daripada mengonsumsi makanan instan atau ultra-processed food."),
        ("Kesadaran Konsumsi Berkelanjutan", "Visualisasi metrik kg terselamatkan di tab profil membangun kebiasaan mindful consumption jangka panjang.")
    ]

    for title, desc in imp_items:
        p_t = tf_i.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.bold = True
        p_t.font.size = Pt(11.5)
        p_t.font.color.rgb = C_TEAL_DARK
        
        r = p_t.add_run()
        r.text = desc
        r.font.bold = False
        r.font.color.rgb = C_TEXT_MUTED
        p_t.space_after = Pt(8)

    # Right: Skalabilitas & Roadmap
    create_card(s9, 6.85, 1.8, 5.65, 4.8)
    tb_scale = s9.shapes.add_textbox(Inches(7.15), Inches(2.1), Inches(5.05), Inches(4.2))
    tf_s = tb_scale.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "Peta Jalan Skalabilitas (Roadmap)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(12)

    scale_items = [
        ("Fase 1 (Selesai): Core App & AI Pipeline", "Deteksi visual Gemini, kurasi Langflow V2, autentikasi multi-provider, dan build release universal."),
        ("Fase 2: Auto-Deduct Inventory Pantry", "Stok bahan di profil otomatis berkurang setelah pengguna menyelesaikan panduan memasak."),
        ("Fase 3: Smart Grocery & Expiry Reminder", "Notifikasi bahan yang mendekati masa kedaluwarsa serta rekomendasi daftar belanja hemat otomatis."),
        ("Fase 4: Komunitas & Food Sharing", "Integrasi donasi bahan makanan berlebih kepada tetangga atau bank makanan lokal.")
    ]

    for title, desc in scale_items:
        p_t = tf_s.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.bold = True
        p_t.font.size = Pt(11.5)
        p_t.font.color.rgb = C_EMERALD
        
        r = p_t.add_run()
        r.text = desc
        r.font.bold = False
        r.font.color.rgb = C_TEXT_MUTED
        p_t.space_after = Pt(8)

    # ==========================================
    # SLIDE 10: CLOSING & LINKS
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10, C_BG_DARK)

    # Accent decorative block
    accent = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(5.0))
    accent.fill.solid()
    accent.fill.fore_color.rgb = C_EMERALD
    accent.line.fill.background()

    tb = s10.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(11), Inches(5.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "TERIMA KASIH — NOMZ AI FOOD RESCUE ASSISTANT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = C_EMERALD
    p0.space_after = Pt(12)

    p1 = tf.add_paragraph()
    p1.text = "Ubah Bahan Sisa Jadi Masakan Berharga."
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Nomz hadir untuk menjawab masalah nyata food waste dengan pendekatan AI yang bertanggung jawab, transparan, dan mudah digunakan oleh semua kalangan."
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "Tautan Proyek & Akses Publik:"
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = C_EMERALD
    p3.space_after = Pt(8)

    links = [
        "GitHub Repository: https://github.com/subjectcode/NomzAI",
        "Official Release v1.0.0: https://github.com/subjectcode/NomzAI/releases/tag/v1.0.0",
        "Direct APK Download: https://github.com/subjectcode/NomzAI/releases/download/v1.0.0/Nomz-v1.0.0-android-universal.apk",
        "Production API: https://nomz-backend-production.up.railway.app"
    ]

    for lnk in links:
        p_lnk = tf.add_paragraph()
        p_lnk.text = f"→ {lnk}"
        p_lnk.font.size = Pt(13)
        p_lnk.font.color.rgb = C_WHITE
        p_lnk.space_after = Pt(4)

    # Save presentation
    output_path = os.path.join(base_dir, "submission", "Nomz_Pitching_Deck.pptx")
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

    # Check file size
    fsize = os.path.getsize(output_path)
    print(f"File size: {fsize:,} bytes ({fsize / (1024*1024):.2f} MB)")

if __name__ == "__main__":
    create_deck()
